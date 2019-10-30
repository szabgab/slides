import openpyxl

wb = openpyxl.load_workbook(filename = 'chart.xlsx')
for ws in wb.worksheets:
    print(ws.title)

ws = wb.worksheets[0]
c = ["Third", 40, 20, 35, 25, 20, 35]

for i in range(len(c)):
    ws.cell(row=i+1, column=3).value = c[i]

lc = openpyxl.chart.LineChart()
lc.title = "Three Lines Chart"
data = openpyxl.chart .Reference(ws,
                                 min_col=1,
                                 min_row=1,
                                 max_col=3,
                                 max_row=len(c))
lc.add_data(data, titles_from_data=True)

ws.add_chart(lc, "D1")

wb.save("chart.xlsx")
